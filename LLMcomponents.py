import pandas as pd
import os
from docx import Document
import re
import openai
import json
import requests
import time


class SettingForLLM():

    def __init__(self):
        self.call_logs = []
        self.last_call_log = None

    def _normalize_usage(self, usage):
        if not usage:
            return {
                "prompt_tokens": 0,
                "completion_tokens": 0,
                "total_tokens": 0
            }

        prompt_tokens = usage.get("prompt_tokens", 0) or 0
        completion_tokens = usage.get("completion_tokens", 0) or 0
        total_tokens = usage.get("total_tokens", 0) or (prompt_tokens + completion_tokens)

        return {
            "prompt_tokens": int(prompt_tokens),
            "completion_tokens": int(completion_tokens),
            "total_tokens": int(total_tokens)
        }

    def _record_call_log(
        self,
        call_type,
        model,
        runtime_seconds,
        response_status_code=None,
        usage=None,
        success=True,
        error_message=None
    ):
        log = {
            "call_type": call_type,
            "model": model,
            "runtime_seconds": float(runtime_seconds),
            "response_status_code": response_status_code,
            "usage": self._normalize_usage(usage),
            "success": bool(success),
            "error_message": error_message
        }

        self.last_call_log = log
        self.call_logs.append(log)
        return log

    def get_last_call_log(self):
        return self.last_call_log

    def get_call_logs(self):
        return self.call_logs

    def clear_call_logs(self):
        self.call_logs = []
        self.last_call_log = None

    def save_call_logs_to_json(self, filename="llm_call_logs.json"):
        with open(filename, "w", encoding="utf-8") as f:
            json.dump(self.call_logs, f, ensure_ascii=False, indent=2)

    def set_chatGPT(self, background, key, chatmodel="gpt-3.5-turbo", url="https://api.openai.com/v1/chat/completions"):
        openai.api_key = key
        headers = {
            "Content-Type": "application/json",
            "Authorization": f"Bearer {openai.api_key}"
        }
        background = background
        messages = [{"role": "system", "content": background}]
        return (url, background, chatmodel, headers, messages)

    def set_payload(self, message, GPTmodelname="gpt-3.5-turbo", messages=[]):
        messages.append({"role": "user", "content": message})
        payload = {
            "model": GPTmodelname,
            "messages": messages,
            "temperature": 1.0,
            "top_p": 1.0,
            "n": 1,
            "stream": False,
            "presence_penalty": 0,
            "frequency_penalty": 0,
        }
        return (payload, messages)

    def send_response_receive_output(self, URL, headers, payload, messages):
        start_time = time.time()
        response = requests.post(URL, headers=headers, json=payload, stream=False)
        runtime_seconds = time.time() - start_time

        try:
            result = json.loads(response.content)
            print(result)

            output = result["choices"][0]["message"]["content"]
            usage = result.get("usage", None)

            self._record_call_log(
                call_type="chat_completion",
                model=payload.get("model"),
                runtime_seconds=runtime_seconds,
                response_status_code=response.status_code,
                usage=usage,
                success=True,
                error_message=None
            )

            messages.append({"role": "assistant", "content": output})
            return (output, messages)

        except Exception as e:
            self._record_call_log(
                call_type="chat_completion",
                model=payload.get("model"),
                runtime_seconds=runtime_seconds,
                response_status_code=response.status_code,
                usage=None,
                success=False,
                error_message=str(e)
            )
            raise e

    def question_matching(self, question, content, model, url, headers, messages):
        query = (
            "My question is: " + question + "\n"
            "Please refer to the following question bank and choose the Section number "
            "(for example, if you choose Section 5, please return 5.) that matches the meaning of my question. "
            "Please note that as long as the meaning matches, there is no need for word-for-word correspondence. "
            "My entry may have spelling or grammatical mistakes, please ignore those mistakes. "
            "Returns 0 if no section matches. Only answer an integer as you choose, do not reply with any information "
            "other than the integer, do not reply why you chose the section number, do not reply to your thought process. "
            "Following is the question bank: \n" + content
        )

        print("Sending question matching prompt to OpenAI:")
        print(query)

        messages.append({"role": "user", "content": query})
        payload = {
            "model": model,
            "messages": messages,
            "temperature": 0.2,
            "top_p": 1.0,
            "n": 1,
            "stream": False,
            "presence_penalty": 0,
            "frequency_penalty": 0,
        }

        start_time = time.time()
        response = requests.post(url, headers=headers, json=payload)
        runtime_seconds = time.time() - start_time

        print("Question Matching Response:")
        print(response.content)

        try:
            result = json.loads(response.content)
            output = result["choices"][0]["message"]["content"].strip()
            usage = result.get("usage", None)

            self._record_call_log(
                call_type="question_matching",
                model=model,
                runtime_seconds=runtime_seconds,
                response_status_code=response.status_code,
                usage=usage,
                success=True,
                error_message=None
            )

        except Exception as e:
            output = "0"
            print(f"❌ Failed to parse OpenAI response: {e}")

            self._record_call_log(
                call_type="question_matching",
                model=model,
                runtime_seconds=runtime_seconds,
                response_status_code=response.status_code,
                usage=None,
                success=False,
                error_message=str(e)
            )

        return output

    def save_chat_history_to_docx(self, messages):
        doc = Document()

        for message in messages:
            role = message['role']
            content = message['content']

            if role == 'user':
                doc.add_paragraph('User: ' + content)
            elif role == 'assistant':
                doc.add_paragraph('Assistant: ' + content)
            else:
                doc.add_paragraph(content)

        filename = "data_report.docx"
        doc.save(filename)